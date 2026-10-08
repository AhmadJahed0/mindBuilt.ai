from docx import Document as DocxDocument


def parse_docx(path: str) -> list[dict]:
    """DOCX has no native page concept, so the whole file is treated as one unit."""
    doc = DocxDocument(path)

    parts = [p.text for p in doc.paragraphs if p.text.strip()]

    for table in doc.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            parts.append(" | ".join(cells))

    return [{"page": 1, "text": "\n".join(parts)}]
