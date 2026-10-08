import openpyxl


def parse_xlsx(path: str) -> list[dict]:
    """Returns one entry per sheet, with sheet name kept as context for the chunk."""
    wb = openpyxl.load_workbook(path, data_only=True)
    sheets = []

    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        rows_text = []
        for row in ws.iter_rows(values_only=True):
            cells = [str(c) for c in row if c is not None]
            if cells:
                rows_text.append(" | ".join(cells))

        sheets.append({
            "page": sheet_name,  # sheet_name doubles as the "page" identifier
            "text": f"Sheet: {sheet_name}\n" + "\n".join(rows_text),
        })

    return sheets
